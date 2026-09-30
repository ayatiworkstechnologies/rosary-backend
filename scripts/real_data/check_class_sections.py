from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"

MASTER_FILE = (
    INPUT_DIR
    / "ROSARY STUDENTS DETAILS 2026 (1).xlsx"
)


def clean(value):
    if value is None:
        return None

    value = str(value).strip()

    return value or None


def main():
    workbook = load_workbook(
        MASTER_FILE,
        read_only=True,
        data_only=True,
    )

    sheet = workbook[
        "Student Parent Details"
    ]

    counts = Counter()

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True,
    ):

        # Based on current master structure:
        # B = Admission No
        # F = Class / Grade
        # G = Section

        admission_no = clean(row[1])
        class_name = clean(row[5])
        section = clean(row[6])

        if not admission_no:
            continue

        key = (
            class_name or "MISSING_CLASS",
            section or "MISSING_SECTION",
        )

        counts[key] += 1

    workbook.close()

    print("=" * 60)
    print("MASTER FILE - CLASS / SECTION COVERAGE")
    print("=" * 60)

    total = 0

    for (
        class_name,
        section,
    ), count in sorted(
        counts.items()
    ):

        print(
            f"{class_name:10} "
            f"{section:5} "
            f"-> {count} students"
        )

        total += count

    print()
    print("=" * 60)

    print(
        f"Total students counted: "
        f"{total}"
    )

    print(
        f"Class/section groups: "
        f"{len(counts)}"
    )


if __name__ == "__main__":
    main()