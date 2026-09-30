from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"


STRUCTURED_FILES = [
    INPUT_DIR
    / "ROSARY STUDENTS DETAILS 2026 (1).xlsx",

    INPUT_DIR
    / "4D,5A,5D6C,6D,7C,8A,8B,8C,9D,10D.xlsx",
]


CLASS_FILES = [
    INPUT_DIR / "1 A.xlsx",
    INPUT_DIR / "II A.xlsx",
    INPUT_DIR / "II B.xlsx",
    INPUT_DIR / "II D.xlsx",
    INPUT_DIR / "III A.xlsx",
    INPUT_DIR / "LKG B.xlsx",
]


def clean(value):
    if value is None:
        return None

    value = str(value).strip()

    return value or None


def main():
    classes = set()

    # ------------------------------------------
    # STRUCTURED FILES
    # ------------------------------------------

    for file_path in STRUCTURED_FILES:

        if not file_path.exists():
            print(
                f"WARNING: Missing file: "
                f"{file_path.name}"
            )
            continue

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        if (
            "Student Parent Details"
            not in workbook.sheetnames
        ):
            print(
                f"WARNING: Student Parent Details "
                f"sheet missing in {file_path.name}"
            )

            workbook.close()
            continue

        sheet = workbook[
            "Student Parent Details"
        ]

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            class_name = clean(
                row[5]
            )

            section = clean(
                row[6]
            )

            academic_year = clean(
                row[8]
            )

            if (
                class_name
                and section
                and academic_year
            ):
                classes.add(
                    (
                        class_name,
                        section,
                        academic_year,
                    )
                )

        workbook.close()

    # ------------------------------------------
    # CLASS FILES WITHOUT HEADERS
    # ------------------------------------------

    for file_path in CLASS_FILES:

        if not file_path.exists():
            print(
                f"WARNING: Missing file: "
                f"{file_path.name}"
            )
            continue

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        for sheet in workbook.worksheets:

            for row in sheet.iter_rows(
                values_only=True,
            ):
                if not row:
                    continue

                if len(row) < 8:
                    continue

                class_name = clean(
                    row[4]
                )

                section = clean(
                    row[5]
                )

                academic_year = clean(
                    row[7]
                )

                if (
                    class_name
                    and section
                    and academic_year
                ):
                    classes.add(
                        (
                            class_name,
                            section,
                            academic_year,
                        )
                    )

        workbook.close()

    # ------------------------------------------
    # OUTPUT
    # ------------------------------------------

    print("=" * 60)
    print("REQUIRED SCHOOL CLASSES")
    print("=" * 60)

    for (
        class_name,
        section,
        academic_year,
    ) in sorted(classes):

        print(
            f"{class_name:<6} "
            f"{section:<5} "
            f"{academic_year}"
        )

    print()

    print(
        f"Total class/section combinations: "
        f"{len(classes)}"
    )


if __name__ == "__main__":
    main()