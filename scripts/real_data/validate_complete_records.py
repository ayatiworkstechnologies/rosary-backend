from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
import re

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"

STRUCTURED_FILES = [
    INPUT_DIR
    / "ROSARY STUDENTS DETAILS 2026 (1).xlsx",

    INPUT_DIR
    / "4D,5A,5D6C,6D,7C,8A,8B,8C,9D,10D.xlsx",
]

CLASS_FILES = [
    INPUT_DIR / "1 A.xlsx",
    INPUT_DIR / "II A.xlsx",
    INPUT_DIR / "II B.xlsx",
    INPUT_DIR / "II D.xlsx",
    INPUT_DIR / "III A.xlsx",
    INPUT_DIR / "LKG B.xlsx",
]


# Fields we actually need for the current
# Student + Parent import.
REQUIRED_FIELDS = [
    "admission_no",
    "student_name",
    "class_name",
    "section",
    "academic_year",
]


def clean(value):
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()

        if not value:
            return None

    return value


def clean_text(value):
    value = clean(value)

    if value is None:
        return None

    return str(value).strip()


def clean_admission_no(value):
    value = clean_text(value)

    if not value:
        return None

    if value.endswith(".0"):
        value = value[:-2]

    return value


def clean_phone(value):
    value = clean(value)

    if value is None:
        return None

    # Excel may give phone as int/float.
    if isinstance(value, float):
        if value.is_integer():
            value = int(value)

    value = str(value).strip()

    if value.endswith(".0"):
        value = value[:-2]

    # Remove spaces, +91, hyphens, brackets, etc.
    digits = re.sub(
        r"\D",
        "",
        value,
    )

    if digits.startswith("91") and len(digits) == 12:
        digits = digits[2:]

    return digits or None


def valid_phone(value):
    phone = clean_phone(value)

    if not phone:
        return False

    return (
        len(phone) == 10
        and phone.isdigit()
    )


def valid_email(value):
    value = clean_text(value)

    if not value:
        return False

    pattern = (
        r"^[A-Za-z0-9._%+-]+"
        r"@[A-Za-z0-9.-]+"
        r"\.[A-Za-z]{2,}$"
    )

    return bool(
        re.match(
            pattern,
            value,
        )
    )


def parse_date(value):
    value = clean(value)

    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    value = str(value).strip()

    formats = [
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%Y/%m/%d",
    ]

    for date_format in formats:
        try:
            return datetime.strptime(
                value,
                date_format,
            ).date()

        except ValueError:
            continue

    return None


def normalize_pincode(value):
    value = clean(value)

    if value is None:
        return None

    if isinstance(value, float):
        if value.is_integer():
            value = int(value)

    value = str(value).strip()

    if value.endswith(".0"):
        value = value[:-2]

    return value


def valid_pincode(value):
    value = normalize_pincode(
        value
    )

    if not value:
        return False

    return (
        value.isdigit()
        and len(value) == 6
    )


def build_record(
    source_file,
    source_sheet,
    source_row,
    admission_no,
    student_name,
    date_of_birth,
    gender,
    class_name,
    section,
    roll_no,
    academic_year,
    admission_date,
    blood_group,
    address,
    city,
    state,
    pincode,
    father_name,
    father_phone,
    parent_email,
    mother_name,
    mother_phone,
):
    return {
        "source_file": source_file,
        "source_sheet": source_sheet,
        "source_row": source_row,

        "admission_no":
            clean_admission_no(
                admission_no
            ),

        "student_name":
            clean_text(
                student_name
            ),

        "date_of_birth":
            parse_date(
                date_of_birth
            ),

        "raw_date_of_birth":
            date_of_birth,

        "gender":
            clean_text(
                gender
            ),

        "class_name":
            clean_text(
                class_name
            ),

        "section":
            clean_text(
                section
            ),

        "roll_no":
            clean_text(
                roll_no
            ),

        "academic_year":
            clean_text(
                academic_year
            ),

        "admission_date":
            parse_date(
                admission_date
            ),

        "blood_group":
            clean_text(
                blood_group
            ),

        "address":
            clean_text(
                address
            ),

        "city":
            clean_text(
                city
            ),

        "state":
            clean_text(
                state
            ),

        "pincode":
            normalize_pincode(
                pincode
            ),

        "father_name":
            clean_text(
                father_name
            ),

        "father_phone":
            clean_phone(
                father_phone
            ),

        "parent_email":
            clean_text(
                parent_email
            ),

        "mother_name":
            clean_text(
                mother_name
            ),

        "mother_phone":
            clean_phone(
                mother_phone
            ),
    }


def read_structured_files():
    records = []

    for file_path in STRUCTURED_FILES:

        if not file_path.exists():
            print(
                f"WARNING: Missing file: "
                f"{file_path.name}"
            )
            continue

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        if (
            "Student Parent Details"
            not in workbook.sheetnames
        ):
            print(
                f"WARNING: Student Parent Details "
                f"sheet missing in {file_path.name}"
            )

            workbook.close()
            continue

        sheet = workbook[
            "Student Parent Details"
        ]

        # Row 1 is the header.
        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):
            if not any(
                value is not None
                for value in row
            ):
                continue

            record = build_record(
                source_file=file_path.name,
                source_sheet=sheet.title,
                source_row=row_number,

                admission_no=row[1],
                student_name=row[2],
                date_of_birth=row[3],
                gender=row[4],
                class_name=row[5],
                section=row[6],
                roll_no=row[7],
                academic_year=row[8],
                admission_date=row[9],
                blood_group=row[10],
                address=row[11],
                city=row[12],
                state=row[13],
                pincode=row[14],
                father_name=row[15],
                father_phone=row[16],
                parent_email=row[17],
                mother_name=row[18],
                mother_phone=row[19],
            )

            records.append(
                record
            )

        workbook.close()

    return records
    records = []

    workbook = load_workbook(
        MASTER_FILE,
        read_only=True,
        data_only=True,
    )

    sheet = workbook[
        "Student Parent Details"
    ]

    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):
        if not any(
            value is not None
            for value in row
        ):
            continue

        record = build_record(
            source_file=MASTER_FILE.name,
            source_sheet=sheet.title,
            source_row=row_number,

            admission_no=row[1],
            student_name=row[2],
            date_of_birth=row[3],
            gender=row[4],
            class_name=row[5],
            section=row[6],
            roll_no=row[7],
            academic_year=row[8],
            admission_date=row[9],
            blood_group=row[10],
            address=row[11],
            city=row[12],
            state=row[13],
            pincode=row[14],
            father_name=row[15],
            father_phone=row[16],
            parent_email=row[17],
            mother_name=row[18],
            mother_phone=row[19],
        )

        records.append(
            record
        )

    workbook.close()

    return records


def read_class_files():
    records = []

    for file_path in CLASS_FILES:

        if not file_path.exists():
            print(
                f"WARNING: Missing file: "
                f"{file_path.name}"
            )

            continue

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        for sheet in workbook.worksheets:

            for row_number, row in enumerate(
                sheet.iter_rows(
                    values_only=True,
                ),
                start=1,
            ):
                if not any(
                    value is not None
                    for value in row
                ):
                    continue

                # Separate class files have no header.
                #
                # 0  Admission No
                # 1  Student Name
                # 2  DOB
                # 3  Gender
                # 4  Class
                # 5  Section
                # 6  Roll No
                # 7  Academic Year
                # 8  Admission Date
                # 9  Blood Group
                # 10 Address
                # 11 City
                # 12 State
                # 13 Pincode
                # 14 Father Name
                # 15 Father Phone
                # 16 Parent Email
                # 17 Mother Name
                # 18 Mother Phone

                if len(row) < 19:
                    continue

                record = build_record(
                    source_file=file_path.name,
                    source_sheet=sheet.title,
                    source_row=row_number,

                    admission_no=row[0],
                    student_name=row[1],
                    date_of_birth=row[2],
                    gender=row[3],
                    class_name=row[4],
                    section=row[5],
                    roll_no=row[6],
                    academic_year=row[7],
                    admission_date=row[8],
                    blood_group=row[9],
                    address=row[10],
                    city=row[11],
                    state=row[12],
                    pincode=row[13],
                    father_name=row[14],
                    father_phone=row[15],
                    parent_email=row[16],
                    mother_name=row[17],
                    mother_phone=row[18],
                )

                records.append(
                    record
                )

        workbook.close()

    return records


def identity_key(record):
    return (
        (
            record["student_name"]
            or ""
        ).upper(),

        (
            record["class_name"]
            or ""
        ).upper(),

        (
            record["section"]
            or ""
        ).upper(),

        str(
            record["roll_no"]
            or ""
        ).upper(),
    )


def validate_record(record):
    errors = []

    # Only core student fields block import.
    for field in REQUIRED_FIELDS:
        if record.get(field) in (
            None,
            "",
        ):
            errors.append(
                f"Missing {field}"
            )

    return errors

def main():
    master_records = (
    read_structured_files()
)

    class_records = (
        read_class_files()
    )

    all_records = (
        master_records
        + class_records
    )

    print("=" * 75)
    print("ROSARY STUDENT + PARENT DATA VALIDATION")
    print("=" * 75)

    print(
        f"Master rows           : "
        f"{len(master_records)}"
    )

    print(
        f"Class-file rows       : "
        f"{len(class_records)}"
    )

    print(
        f"Total source rows     : "
        f"{len(all_records)}"
    )

    # --------------------------------------------------
    # GROUP BY ADMISSION NUMBER
    # --------------------------------------------------

    grouped = defaultdict(list)

    missing_admission = []

    for record in all_records:

        admission_no = (
            record["admission_no"]
        )

        if not admission_no:
            missing_admission.append(
                record
            )

            continue

        grouped[
            admission_no
        ].append(
            record
        )

    # --------------------------------------------------
    # HANDLE DUPLICATES
    # --------------------------------------------------

    resolved_records = []
    duplicate_same_student = []
    conflicting_admissions = []

    for admission_no, records in (
        grouped.items()
    ):

        if len(records) == 1:
            resolved_records.append(
                records[0]
            )

            continue

        identities = {
            identity_key(record)
            for record in records
        }

        if len(identities) == 1:

            # Same student appears more than once.
            # Prefer master record if available.

            master_matches = [
                record
                for record in records
                if record["source_file"] in {
                    file_path.name
                    for file_path in STRUCTURED_FILES
                }
            ]

            chosen = (
                master_matches[0]
                if master_matches
                else records[0]
            )

            resolved_records.append(
                chosen
            )

            duplicate_same_student.append(
                (
                    admission_no,
                    records,
                )
            )

        else:
            # Same admission number belongs to
            # different students.
            # Do not import any of them.

            conflicting_admissions.append(
                (
                    admission_no,
                    records,
                )
            )

    # --------------------------------------------------
    # VALIDATE COMPLETE RECORDS
    # --------------------------------------------------

    valid_records = []
    invalid_records = []

    for record in resolved_records:

        errors = validate_record(
            record
        )

        if errors:
            invalid_records.append(
                (
                    record,
                    errors,
                )
            )
        else:
            valid_records.append(
                record
            )

    # --------------------------------------------------
    # SUMMARY
    # --------------------------------------------------

    print()
    print("=" * 75)
    print("SUMMARY")
    print("=" * 75)

    print(
        f"Unique resolved students      : "
        f"{len(resolved_records)}"
    )

    print(
        f"Same-student duplicates       : "
        f"{len(duplicate_same_student)}"
    )

    print(
        f"Conflicting admission numbers : "
        f"{len(conflicting_admissions)}"
    )

    print(
        f"Rows missing admission no     : "
        f"{len(missing_admission)}"
    )

    print()
    print(
        f"COMPLETE / IMPORT-READY       : "
        f"{len(valid_records)}"
    )

    print(
        f"INCOMPLETE / SKIPPED          : "
        f"{len(invalid_records)}"
    )

    # --------------------------------------------------
    # CONFLICT DETAILS
    # --------------------------------------------------

    if conflicting_admissions:

        print()
        print("=" * 75)
        print(
            "CONFLICTING ADMISSION NUMBERS "
            "- WILL NOT IMPORT"
        )
        print("=" * 75)

        for (
            admission_no,
            records,
        ) in conflicting_admissions:

            print()
            print(
                f"Admission No: "
                f"{admission_no}"
            )

            for record in records:

                print(
                    f"  {record['student_name']} "
                    f"| "
                    f"{record['class_name']}-"
                    f"{record['section']} "
                    f"| "
                    f"{record['source_file']} "
                    f"| Row "
                    f"{record['source_row']}"
                )

    # --------------------------------------------------
    # INVALID DETAILS
    # --------------------------------------------------

    print()
    print("=" * 75)
    print("INCOMPLETE / INVALID RECORDS")
    print("=" * 75)

    if not invalid_records:
        print(
            "No invalid records."
        )

    else:
        for record, errors in (
            invalid_records
        ):

            print()
            print(
                f"Admission No : "
                f"{record['admission_no']}"
            )

            print(
                f"Student      : "
                f"{record['student_name']}"
            )

            print(
                f"Class        : "
                f"{record['class_name']}"
                f"-"
                f"{record['section']}"
            )

            print(
                f"Source       : "
                f"{record['source_file']} "
                f"Row "
                f"{record['source_row']}"
            )

            print(
                "Problems:"
            )

            for error in errors:
                print(
                    f"  - {error}"
                )


if __name__ == "__main__":
    main()