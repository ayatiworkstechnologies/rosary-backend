from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"


MASTER_FILE = (
    INPUT_DIR
    / "ROSARY STUDENTS DETAILS 2026 (1).xlsx"
)


CLASS_FILES = [
    INPUT_DIR / "1 A.xlsx",
    INPUT_DIR / "II A.xlsx",
    INPUT_DIR / "II B.xlsx",
    INPUT_DIR / "II D.xlsx",
    INPUT_DIR / "III A.xlsx",
    INPUT_DIR / "LKG B.xlsx",
]


def clean_admission_no(value):
    if value is None:
        return None

    value = str(value).strip()

    if value.endswith(".0"):
        value = value[:-2]

    return value or None


def get_master_admission_numbers():
    workbook = load_workbook(
        MASTER_FILE,
        read_only=True,
        data_only=True,
    )

    sheet = workbook[
        "Student Parent Details"
    ]

    admission_numbers = set()

    # Row 1 = header
    # Admission No is column B
    for row in sheet.iter_rows(
        min_row=2,
        values_only=True,
    ):
        admission_no = (
            clean_admission_no(
                row[1]
            )
        )

        if admission_no:
            admission_numbers.add(
                admission_no
            )

    workbook.close()

    return admission_numbers


def get_class_file_admission_numbers(
    file_path,
):
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    admission_numbers = set()

    for sheet in workbook.worksheets:

        # These files have NO header.
        # Column 1 = Admission No.
        for row in sheet.iter_rows(
            values_only=True,
        ):
            if not row:
                continue

            admission_no = (
                clean_admission_no(
                    row[0]
                )
            )

            if admission_no:
                admission_numbers.add(
                    admission_no
                )

    workbook.close()

    return admission_numbers


def main():

    master_numbers = (
        get_master_admission_numbers()
    )

    print(
        f"Master students: "
        f"{len(master_numbers)}"
    )

    print()

    all_class_numbers = set()

    for file_path in CLASS_FILES:

        if not file_path.exists():
            print(
                f"Missing file: "
                f"{file_path.name}"
            )
            continue

        class_numbers = (
            get_class_file_admission_numbers(
                file_path
            )
        )

        all_class_numbers.update(
            class_numbers
        )

        overlapping = (
            class_numbers
            & master_numbers
        )

        only_in_class = (
            class_numbers
            - master_numbers
        )

        print("=" * 60)

        print(
            f"FILE: {file_path.name}"
        )

        print(
            f"Students: "
            f"{len(class_numbers)}"
        )

        print(
            f"Already in master: "
            f"{len(overlapping)}"
        )

        print(
            f"Only in class file: "
            f"{len(only_in_class)}"
        )

        if only_in_class:

            print(
                "Admission numbers "
                "not in master:"
            )

            for admission_no in sorted(
                only_in_class
            ):
                print(
                    f"  - {admission_no}"
                )

    # --------------------------------------------------
    # STUDENTS PRESENT IN BOTH SOURCES
    # --------------------------------------------------

    common_students = (
        all_class_numbers
        & master_numbers
    )

    print()
    print("=" * 60)
    print("STUDENTS PRESENT IN BOTH SOURCES")
    print("=" * 60)

    if common_students:

        for admission_no in sorted(
            common_students
        ):
            print(
                f"  - {admission_no}"
            )

    else:
        print(
            "No overlapping students found."
        )

    # --------------------------------------------------
    # OVERALL SUMMARY
    # --------------------------------------------------

    print()
    print("=" * 60)
    print("OVERALL")
    print("=" * 60)

    print(
        f"Unique students in "
        f"class files: "
        f"{len(all_class_numbers)}"
    )

    print(
        f"Class students also "
        f"in master: "
        f"{len(common_students)}"
    )

    print(
        f"Class students NOT "
        f"in master: "
        f"{len(all_class_numbers - master_numbers)}"
    )


if __name__ == "__main__":
    main()