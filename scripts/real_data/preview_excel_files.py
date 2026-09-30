from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"


def main():
    excel_files = sorted(
        INPUT_DIR.glob("*.xlsx")
    )

    for file_path in excel_files:

        print("\n")
        print("=" * 80)
        print(f"FILE: {file_path.name}")
        print("=" * 80)

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        for worksheet in workbook.worksheets:

            print(
                f"\nSHEET: "
                f"{worksheet.title}"
            )

            print("-" * 80)

            for row_number, row in enumerate(
                worksheet.iter_rows(
                    values_only=True
                ),
                start=1,
            ):

                print(
                    f"ROW {row_number}: "
                    f"{row}"
                )

                if row_number >= 3:
                    break

        workbook.close()


if __name__ == "__main__":
    main()