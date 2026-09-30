from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"

MASTER_FILE = (
    INPUT_DIR
    / "ROSARY STUDENTS DETAILS 2026 (1).xlsx"
)


def clean(value):
    if value is None:
        return None

    value = str(value).strip()

    if value.endswith(".0"):
        value = value[:-2]

    return value or None


def main():
    workbook = load_workbook(
        MASTER_FILE,
        read_only=True,
        data_only=True,
    )

    sheet = workbook[
        "Student Parent Details"
    ]

    students = defaultdict(list)

    for excel_row, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):
        admission_no = clean(row[1])

        if not admission_no:
            continue

        students[admission_no].append(
            {
                "excel_row": excel_row,
                "name": clean(row[2]),
                "class": clean(row[5]),
                "section": clean(row[6]),
                "roll_no": clean(row[7]),
                "father_name": clean(row[15]),
                "mother_name": clean(row[18]),
            }
        )

    workbook.close()

    duplicates = {
        admission_no: records
        for admission_no, records
        in students.items()
        if len(records) > 1
    }

    print("=" * 70)
    print("DUPLICATE ADMISSION NUMBER CHECK")
    print("=" * 70)

    print(
        f"Duplicate admission numbers: "
        f"{len(duplicates)}"
    )

    for admission_no, records in (
        duplicates.items()
    ):
        print()
        print(
            f"Admission No: "
            f"{admission_no}"
        )

        for record in records:
            print(
                f"  Excel Row : "
                f"{record['excel_row']}"
            )

            print(
                f"  Student   : "
                f"{record['name']}"
            )

            print(
                f"  Class     : "
                f"{record['class']} - "
                f"{record['section']}"
            )

            print(
                f"  Roll No   : "
                f"{record['roll_no']}"
            )

            print(
                f"  Father    : "
                f"{record['father_name']}"
            )

            print(
                f"  Mother    : "
                f"{record['mother_name']}"
            )

            print("  " + "-" * 40)


if __name__ == "__main__":
    main()