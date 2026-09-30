from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"

TARGET_ADMISSION_NO = "19332"

MASTER_FILE = (
    INPUT_DIR
    / "ROSARY STUDENTS DETAILS 2026 (1).xlsx"
)

CLASS_FILES = [
    INPUT_DIR / "1 A.xlsx",
    INPUT_DIR / "II A.xlsx",
    INPUT_DIR / "II B.xlsx",
    INPUT_DIR / "II D.xlsx",
    INPUT_DIR / "III A.xlsx",
    INPUT_DIR / "LKG B.xlsx",
]


def clean(value):
    if value is None:
        return None

    value = str(value).strip()

    if value.endswith(".0"):
        value = value[:-2]

    return value or None


def print_record(source, values):
    print()
    print("=" * 70)
    print(f"SOURCE: {source}")
    print("=" * 70)

    for key, value in values.items():
        print(
            f"{key:<20}: {value}"
        )


def check_master():
    workbook = load_workbook(
        MASTER_FILE,
        read_only=True,
        data_only=True,
    )

    sheet = workbook[
        "Student Parent Details"
    ]

    for excel_row, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):
        admission_no = clean(row[1])

        if admission_no != TARGET_ADMISSION_NO:
            continue

        print_record(
            f"{MASTER_FILE.name} - Row {excel_row}",
            {
                "Admission No": clean(row[1]),
                "Student Name": clean(row[2]),
                "DOB": row[3],
                "Gender": clean(row[4]),
                "Class": clean(row[5]),
                "Section": clean(row[6]),
                "Roll No": clean(row[7]),
                "Academic Year": clean(row[8]),
                "Father Name": clean(row[15]),
                "Father Phone": clean(row[16]),
                "Parent Email": clean(row[17]),
                "Mother Name": clean(row[18]),
                "Mother Phone": clean(row[19]),
            },
        )

    workbook.close()


def check_class_files():
    for file_path in CLASS_FILES:

        if not file_path.exists():
            continue

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        for sheet in workbook.worksheets:

            for excel_row, row in enumerate(
                sheet.iter_rows(
                    values_only=True,
                ),
                start=1,
            ):
                if not row:
                    continue

                admission_no = clean(row[0])

                if admission_no != TARGET_ADMISSION_NO:
                    continue

                print_record(
                    (
                        f"{file_path.name} - "
                        f"{sheet.title} - "
                        f"Row {excel_row}"
                    ),
                    {
                        "Admission No": clean(row[0]),
                        "Student Name": clean(row[1]),
                        "DOB": row[2],
                        "Gender": clean(row[3]),
                        "Class": clean(row[4]),
                        "Section": clean(row[5]),
                        "Roll No": clean(row[6]),
                        "Academic Year": clean(row[7]),
                        "Father Name": clean(row[14]),
                        "Father Phone": clean(row[15]),
                        "Parent Email": clean(row[16]),
                        "Mother Name": clean(row[17]),
                        "Mother Phone": clean(row[18]),
                    },
                )

        workbook.close()


def main():
    print(
        f"Checking Admission No: "
        f"{TARGET_ADMISSION_NO}"
    )

    check_master()
    check_class_files()


if __name__ == "__main__":
    main()