from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent

INPUT_DIR = BASE_DIR / "input"


def clean_value(value):
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()

        if not value:
            return None

    return value


def inspect_sheet(
    file_name: str,
    sheet_name: str,
    worksheet,
):
    rows = list(
        worksheet.iter_rows(
            values_only=True
        )
    )

    if not rows:
        print(
            f"\nSheet: {sheet_name}"
        )
        print("Empty sheet")
        return

    # -----------------------------------------------------
    # HEADERS
    # -----------------------------------------------------

    headers = [
        str(value).strip()
        if value is not None
        else ""
        for value in rows[0]
    ]

    valid_indexes = [
        index
        for index, header
        in enumerate(headers)
        if header
    ]

    valid_headers = [
        headers[index]
        for index in valid_indexes
    ]

    # -----------------------------------------------------
    # DATA ROWS
    # -----------------------------------------------------

    data_rows = rows[1:]

    total_rows = 0
    complete_rows = 0
    incomplete_rows = 0

    missing_counts = {
        header: 0
        for header in valid_headers
    }

    for row in data_rows:

        values = [
            clean_value(
                row[index]
                if index < len(row)
                else None
            )
            for index in valid_indexes
        ]

        # Completely blank row
        if all(
            value is None
            for value in values
        ):
            continue

        total_rows += 1

        missing_headers = []

        for header, value in zip(
            valid_headers,
            values,
        ):
            if value is None:
                missing_counts[
                    header
                ] += 1

                missing_headers.append(
                    header
                )

        if missing_headers:
            incomplete_rows += 1
        else:
            complete_rows += 1

    # -----------------------------------------------------
    # REPORT
    # -----------------------------------------------------

    print("\n")
    print("=" * 70)

    print(
        f"FILE  : {file_name}"
    )

    print(
        f"SHEET : {sheet_name}"
    )

    print("=" * 70)

    print(
        f"Columns         : "
        f"{len(valid_headers)}"
    )

    print(
        f"Data rows       : "
        f"{total_rows}"
    )

    print(
        f"Complete rows   : "
        f"{complete_rows}"
    )

    print(
        f"Incomplete rows : "
        f"{incomplete_rows}"
    )

    print("\nColumns:")

    for header in valid_headers:
        print(
            f"  - {header}"
        )

    print("\nMissing values by column:")

    missing_found = False

    for header, count in (
        missing_counts.items()
    ):
        if count > 0:
            missing_found = True

            print(
                f"  {header}: {count}"
            )

    if not missing_found:
        print(
            "  No missing values"
        )


def inspect_file(
    file_path: Path,
):
    try:
        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True,
        )

        for worksheet in (
            workbook.worksheets
        ):
            inspect_sheet(
                file_name=
                    file_path.name,

                sheet_name=
                    worksheet.title,

                worksheet=
                    worksheet,
            )

        workbook.close()

    except Exception as exc:
        print("\n")
        print("=" * 70)

        print(
            f"ERROR READING: "
            f"{file_path.name}"
        )

        print(
            f"Reason: {exc}"
        )


def main():

    if not INPUT_DIR.exists():
        print(
            f"Input folder not found: "
            f"{INPUT_DIR}"
        )

        return

    excel_files = sorted(
        list(
            INPUT_DIR.glob(
                "*.xlsx"
            )
        )
        +
        list(
            INPUT_DIR.glob(
                "*.xlsm"
            )
        )
    )

    if not excel_files:
        print(
            "No Excel files found "
            "inside input folder."
        )

        return

    print(
        f"Found {len(excel_files)} "
        f"Excel file(s)."
    )

    for file_path in excel_files:
        inspect_file(
            file_path
        )


if __name__ == "__main__":
    main()