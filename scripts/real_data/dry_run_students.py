from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
import sys

from openpyxl import load_workbook


# ---------------------------------------------------------
# PROJECT PATH
# ---------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = BASE_DIR.parent.parent

if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(
        0,
        str(PROJECT_ROOT),
    )


from app.database import SessionLocal
from app.models.school_class import SchoolClass
from app.models.student import Student


INPUT_DIR = BASE_DIR / "input"


STRUCTURED_FILES = [
    INPUT_DIR
    / "ROSARY STUDENTS DETAILS 2026 (1).xlsx",

    INPUT_DIR
    / "4D,5A,5D6C,6D,7C,8A,8B,8C,9D,10D.xlsx",
]


CLASS_FILES = [
    INPUT_DIR / "1 A.xlsx",
    INPUT_DIR / "II A.xlsx",
    INPUT_DIR / "II B.xlsx",
    INPUT_DIR / "II D.xlsx",
    INPUT_DIR / "III A.xlsx",
    INPUT_DIR / "LKG B.xlsx",
]


# ---------------------------------------------------------
# CLEANING HELPERS
# ---------------------------------------------------------

def clean_text(value):
    if value is None:
        return None

    if isinstance(value, float):
        if value.is_integer():
            value = int(value)

    value = str(value).strip()

    if value.endswith(".0"):
        value = value[:-2]

    return value or None


def clean_admission_no(value):
    return clean_text(value)


def parse_date(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    value = str(value).strip()

    formats = [
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%Y/%m/%d",
    ]

    for date_format in formats:

        try:
            return datetime.strptime(
                value,
                date_format,
            ).date()

        except ValueError:
            continue

    # Invalid / unusable DOB is allowed.
    # Student will still be imported with NULL DOB.
    return None


def normalize_class_value(value):
    value = clean_text(value)

    if not value:
        return None

    return value.upper()


# ---------------------------------------------------------
# RECORD BUILDER
# ---------------------------------------------------------

def build_student_record(
    source_file,
    source_sheet,
    source_row,
    admission_no,
    full_name,
    date_of_birth,
    gender,
    class_name,
    section,
    roll_no,
    academic_year,
):
    return {
        "source_file": source_file,
        "source_sheet": source_sheet,
        "source_row": source_row,

        "admission_no":
            clean_admission_no(
                admission_no
            ),

        "full_name":
            clean_text(
                full_name
            ),

        "date_of_birth":
            parse_date(
                date_of_birth
            ),

        "gender":
            clean_text(
                gender
            ),

        "class_name":
            normalize_class_value(
                class_name
            ),

        "section":
            normalize_class_value(
                section
            ),

        "roll_no":
            clean_text(
                roll_no
            ),

        "academic_year":
            clean_text(
                academic_year
            ),
    }


# ---------------------------------------------------------
# READ STRUCTURED FILES
# ---------------------------------------------------------

def read_structured_files():
    records = []

    for file_path in STRUCTURED_FILES:

        if not file_path.exists():
            print(
                f"WARNING: File missing: "
                f"{file_path.name}"
            )
            continue

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        if (
            "Student Parent Details"
            not in workbook.sheetnames
        ):
            print(
                f"WARNING: Student Parent Details "
                f"sheet missing in "
                f"{file_path.name}"
            )

            workbook.close()
            continue

        sheet = workbook[
            "Student Parent Details"
        ]

        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):

            if not any(
                value is not None
                for value in row
            ):
                continue

            if len(row) < 9:
                continue

            record = build_student_record(
                source_file=file_path.name,
                source_sheet=sheet.title,
                source_row=row_number,

                admission_no=row[1],
                full_name=row[2],
                date_of_birth=row[3],
                gender=row[4],
                class_name=row[5],
                section=row[6],
                roll_no=row[7],
                academic_year=row[8],
            )

            records.append(
                record
            )

        workbook.close()

    return records


# ---------------------------------------------------------
# READ CLASS FILES WITHOUT HEADERS
# ---------------------------------------------------------

def read_class_files():
    records = []

    for file_path in CLASS_FILES:

        if not file_path.exists():
            print(
                f"WARNING: File missing: "
                f"{file_path.name}"
            )
            continue

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        for sheet in workbook.worksheets:

            for row_number, row in enumerate(
                sheet.iter_rows(
                    values_only=True,
                ),
                start=1,
            ):

                if not row:
                    continue

                if not any(
                    value is not None
                    for value in row
                ):
                    continue

                if len(row) < 8:
                    continue

                record = build_student_record(
                    source_file=file_path.name,
                    source_sheet=sheet.title,
                    source_row=row_number,

                    admission_no=row[0],
                    full_name=row[1],
                    date_of_birth=row[2],
                    gender=row[3],
                    class_name=row[4],
                    section=row[5],
                    roll_no=row[6],
                    academic_year=row[7],
                )

                records.append(
                    record
                )

        workbook.close()

    return records


# ---------------------------------------------------------
# DUPLICATE RESOLUTION
# ---------------------------------------------------------

def identity_key(record):
    return (
        (
            record["full_name"]
            or ""
        ).strip().upper(),

        (
            record["class_name"]
            or ""
        ).strip().upper(),

        (
            record["section"]
            or ""
        ).strip().upper(),

        (
            record["roll_no"]
            or ""
        ).strip().upper(),
    )


def record_quality(record):
    score = 0

    if record["full_name"]:
        score += 1

    if record["date_of_birth"]:
        score += 1

    if record["gender"]:
        score += 1

    if record["roll_no"]:
        score += 1

    if record["class_name"]:
        score += 1

    if record["section"]:
        score += 1

    if record["academic_year"]:
        score += 1

    return score


def resolve_students(records):
    grouped = defaultdict(list)

    missing_admission = []

    for record in records:

        admission_no = (
            record["admission_no"]
        )

        if not admission_no:
            missing_admission.append(
                record
            )
            continue

        grouped[
            admission_no
        ].append(
            record
        )

    resolved = []
    same_student_duplicates = []
    conflicts = []

    for admission_no, matches in (
        grouped.items()
    ):

        if len(matches) == 1:
            resolved.append(
                matches[0]
            )
            continue

        identities = {
            identity_key(record)
            for record in matches
        }

        if len(identities) == 1:

            # Same student appears in multiple files.
            # Keep the most complete student record.

            chosen = max(
                matches,
                key=record_quality,
            )

            resolved.append(
                chosen
            )

            same_student_duplicates.append(
                (
                    admission_no,
                    matches,
                )
            )

        else:

            # Same admission number belongs to
            # different students.
            # Never import automatically.

            conflicts.append(
                (
                    admission_no,
                    matches,
                )
            )

    return (
        resolved,
        same_student_duplicates,
        conflicts,
        missing_admission,
    )


# ---------------------------------------------------------
# REQUIRED STUDENT FIELDS
# ---------------------------------------------------------

def validate_student(record):
    errors = []

    required_fields = [
        "admission_no",
        "full_name",
        "class_name",
        "section",
        "academic_year",
    ]

    for field in required_fields:

        if not record.get(field):
            errors.append(
                f"Missing {field}"
            )

    return errors


# ---------------------------------------------------------
# DATABASE DRY RUN
# ---------------------------------------------------------

def main():

    print("=" * 75)
    print("ROSARY REAL STUDENT IMPORT - DRY RUN")
    print("=" * 75)

    structured_records = (
        read_structured_files()
    )

    class_file_records = (
        read_class_files()
    )

    source_records = (
        structured_records
        + class_file_records
    )

    print()
    print(
        f"Structured file rows : "
        f"{len(structured_records)}"
    )

    print(
        f"Class-file rows      : "
        f"{len(class_file_records)}"
    )

    print(
        f"Total source rows    : "
        f"{len(source_records)}"
    )

    (
        resolved_students,
        same_student_duplicates,
        conflicts,
        missing_admission,
    ) = resolve_students(
        source_records
    )

    valid_students = []
    invalid_students = []

    for record in resolved_students:

        errors = validate_student(
            record
        )

        if errors:
            invalid_students.append(
                (
                    record,
                    errors,
                )
            )
        else:
            valid_students.append(
                record
            )

    print()
    print("=" * 75)
    print("EXCEL VALIDATION")
    print("=" * 75)

    print(
        f"Resolved students          : "
        f"{len(resolved_students)}"
    )

    print(
        f"Same-student duplicates    : "
        f"{len(same_student_duplicates)}"
    )

    print(
        f"Conflicting Admission Nos  : "
        f"{len(conflicts)}"
    )

    print(
        f"Missing Admission Nos      : "
        f"{len(missing_admission)}"
    )

    print(
        f"Valid students             : "
        f"{len(valid_students)}"
    )

    print(
        f"Invalid students           : "
        f"{len(invalid_students)}"
    )

    # -----------------------------------------------------
    # DATABASE CHECK
    # -----------------------------------------------------

    db = SessionLocal()

    try:

        school_classes = (
            db.query(
                SchoolClass
            )
            .all()
        )

        class_map = {}
        duplicate_db_classes = []

        for school_class in school_classes:

            key = (
                normalize_class_value(
                    school_class.name
                ),
                normalize_class_value(
                    school_class.section
                ),
                clean_text(
                    school_class.academic_year
                ),
            )

            if key in class_map:

                duplicate_db_classes.append(
                    key
                )

                continue

            class_map[
                key
            ] = school_class.id

        existing_students = (
            db.query(
                Student
            )
            .all()
        )

        existing_admissions = {
            clean_admission_no(
                student.admission_no
            )
            for student
            in existing_students
        }

        ready_to_insert = []
        already_existing = []
        missing_classes = []

        for record in valid_students:

            admission_no = (
                record["admission_no"]
            )

            if (
                admission_no
                in existing_admissions
            ):
                already_existing.append(
                    record
                )
                continue

            class_key = (
                record["class_name"],
                record["section"],
                record["academic_year"],
            )

            class_id = class_map.get(
                class_key
            )

            if not class_id:

                missing_classes.append(
                    (
                        record,
                        class_key,
                    )
                )

                continue

            # Attach class_id only in memory.
            # Nothing is written to database.

            preview_record = {
                **record,
                "class_id": class_id,
            }

            ready_to_insert.append(
                preview_record
            )

        print()
        print("=" * 75)
        print("DATABASE DRY-RUN RESULT")
        print("=" * 75)

        print(
            f"Existing DB students       : "
            f"{len(existing_students)}"
        )

        print(
            f"Already existing admissions: "
            f"{len(already_existing)}"
        )

        print(
            f"Missing class mappings     : "
            f"{len(missing_classes)}"
        )

        print(
            f"Duplicate DB class keys    : "
            f"{len(duplicate_db_classes)}"
        )

        print()
        print(
            f"READY TO INSERT            : "
            f"{len(ready_to_insert)}"
        )

        print()
        print(
            "DATABASE CHANGES           : 0"
        )

        print(
            "MODE                       : DRY RUN ONLY"
        )

        # -------------------------------------------------
        # CONFLICT DETAILS
        # -------------------------------------------------

        if conflicts:

            print()
            print("=" * 75)
            print(
                "CONFLICTING ADMISSION NUMBERS"
            )
            print("=" * 75)

            for (
                admission_no,
                matches,
            ) in conflicts:

                print()
                print(
                    f"Admission No: "
                    f"{admission_no}"
                )

                for record in matches:

                    print(
                        f"  "
                        f"{record['full_name']} "
                        f"| "
                        f"{record['class_name']}-"
                        f"{record['section']} "
                        f"| "
                        f"{record['source_file']} "
                        f"Row "
                        f"{record['source_row']}"
                    )

        # -------------------------------------------------
        # MISSING CLASS DETAILS
        # -------------------------------------------------

        if missing_classes:

            print()
            print("=" * 75)
            print(
                "MISSING CLASS MAPPINGS"
            )
            print("=" * 75)

            shown = set()

            for (
                record,
                class_key,
            ) in missing_classes:

                if class_key in shown:
                    continue

                shown.add(
                    class_key
                )

                print(
                    f"{class_key[0]} - "
                    f"{class_key[1]} "
                    f"({class_key[2]})"
                )

        # -------------------------------------------------
        # SMALL PREVIEW
        # -------------------------------------------------

        print()
        print("=" * 75)
        print("FIRST 10 READY STUDENTS")
        print("=" * 75)

        for record in (
            ready_to_insert[:10]
        ):

            print(
                f"{record['admission_no']} "
                f"| "
                f"{record['full_name']} "
                f"| "
                f"{record['class_name']}-"
                f"{record['section']} "
                f"| class_id="
                f"{record['class_id']}"
            )

    finally:
        db.close()


if __name__ == "__main__":
    main()